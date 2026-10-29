import io
import re
from datetime import date
import pandas as pd
import streamlit as st
st.markdown("""
<style>
[data-testid="stDataFrame"],
[data-testid="stDataEditor"] {
    border: 1px solid #808080 !important;
    border-radius: 0 !important;
}
[data-testid="stDataFrame"] [role="gridcell"],
[data-testid="stDataEditor"] [role="gridcell"],
[data-testid="stDataFrame"] [role="columnheader"],
[data-testid="stDataEditor"] [role="columnheader"] {
    border-right: 1px solid #b7b7b7 !important;
    border-bottom: 1px solid #b7b7b7 !important;
}
</style>
""", unsafe_allow_html=True)
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from pypdf import PdfReader
from PIL import Image

st.set_page_config(page_title="AI Excel Generator", page_icon="📊", layout="wide")

# Theme-like styling
st.markdown("""
<style>
.block-container {padding-top: 1.6rem; max-width: 1200px;}
.hero {padding: 26px; border-radius: 18px; background: linear-gradient(120deg,#173b75,#2877b8); color: white; margin-bottom: 18px;}
.hero h1 {font-size: 2.1rem; margin: 0 0 8px 0;}
.hero p {font-size: 1rem; margin: 0; opacity: .95;}
div[data-testid="stMetric"] {border: 1px solid #e4e9f2; border-radius: 12px; padding: 12px;}
.stButton>button, .stDownloadButton>button {border-radius: 10px; min-height: 2.8rem;}
.small-note {color: #5b6575; font-size: .9rem;}
</style>
<div class="hero">
  <h1>📊 AI Excel Generator</h1>
  <p>ফাইল আপলোড করুন, ডেটা দেখুন, সাজান এবং Excel ফাইল ডাউনলোড করুন।</p>
</div>
""", unsafe_allow_html=True)

with st.sidebar:
    st.header("⚙️ Menu")
    page = st.radio("বিভাগ নির্বাচন করুন", [
        "🏠 Home", "📂 File Converter", "👨‍💼 Employee Details",
        "💰 Salary Calculator", "🗓️ Attendance Sheet"
    ])
    st.divider()
    st.caption("Starter version • AI API এখনো যুক্ত করা হয়নি")

def style_workbook(wb):
    header_fill = PatternFill("solid", fgColor="173B75")
    header_font = Font(color="FFFFFF", bold=True)
    thin_blue = Side(style="thin", color="D8E2F0")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin_blue)
        ws.row_dimensions[1].height = 28
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            ws.column_dimensions[letter].width = min(max(max_len + 3, 12), 32)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = Border(bottom=Side(style="hair", color="E7ECF3"))

def df_to_xlsx(df, sheet_name="Data"):
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name[:31] or "Data")
        style_workbook(writer.book)
    out.seek(0)
    return out.getvalue()

def template_xlsx(kind):
    wb = Workbook()
    ws = wb.active
    if kind == "Employee Details":
        ws.title = "Employee Details"
        ws.append(["Employee ID", "Employee Name", "Department", "Designation", "Phone", "Joining Date"])
        ws.append(["EMP001", "Example Name", "Accounts", "Assistant", "", ""])
    elif kind == "Salary Calculation":
        ws.title = "Salary Calculation"
        ws.append(["Employee ID", "Employee Name", "Basic Salary", "HRA", "Other Allowance", "Deduction", "Gross Salary", "Net Salary"])
        ws.append(["EMP001", "Example Name", 15000, "=C2*0.20", 1000, 500, "=C2+D2+E2", "=G2-F2"])
        ws.append(["EMP002", "", 0, "=C3*0.20", 0, 0, "=C3+D3+E3", "=G3-F3"])
    else:
        ws.title = "Attendance"
        ws.append(["Employee ID", "Employee Name", "Month", "Working Days", "Present Days", "Absent Days", "Attendance %"])
        ws.append(["EMP001", "Example Name", date.today().strftime("%B %Y"), 26, 25, "=D2-E2", '=IF(D2=0,0,E2/D2)'])
        ws["G2"].number_format = "0.0%"
    style_workbook(wb)
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()

if page == "🏠 Home":
    st.subheader("আপনার কাজ শুরু করুন")
    c1, c2, c3 = st.columns(3)
    with c1:
        st.metric("File Converter", "PDF / Excel")
        st.caption("Text-based PDF বা Excel থেকে ডেটা দেখুন")
    with c2:
        st.metric("Ready Templates", "3")
        st.caption("Employee, Salary, Attendance")
    with c3:
        st.metric("Download", "XLSX")
        st.caption("সাজানো Excel ফাইল ডাউনলোড")
    st.markdown("### কীভাবে ব্যবহার করবেন")
    st.markdown("""
    1. বাম পাশের Menu থেকে প্রয়োজনীয় বিভাগ বেছে নিন।
    2. নিজের ফাইল আপলোড করুন অথবা একটি Template নির্বাচন করুন।
    3. ডেটা যাচাই/এডিট করুন।
    4. **Download Excel** চাপুন।
    """)
    st.warning("এটি প্রাথমিক ওয়েবসাইটের সংস্করণ। এখনো AI দিয়ে নিজে থেকে সব ধরনের ফাইল বুঝে কলাম সাজানোর সুবিধা যুক্ত হয়নি। গুরুত্বপূর্ণ তথ্য ডাউনলোডের আগে যাচাই করুন।")
    st.info("ব্যক্তিগত তথ্যযুক্ত ফাইল আপলোড করার আগে গোপনীয়তার কথা বিবেচনা করুন।")

elif page == "📂 File Converter":
    st.subheader("📂 File Converter")
    st.write("Excel (.xlsx/.xls) অথবা text-based PDF আপলোড করুন।")
    uploaded = st.file_uploader("ফাইল নির্বাচন করুন", type=["xlsx", "xls", "csv", "pdf"])
    if uploaded:
        name = uploaded.name.lower()
        try:
            if name.endswith(".xlsx"):
                df = pd.read_excel(uploaded, engine="openpyxl")
                st.success("Excel ফাইল পড়া হয়েছে। নিচে ডেটা এডিট করতে পারবেন।")
                edited = st.data_editor(df, num_rows="dynamic", use_container_width=True, key="xlsx_editor")
                st.download_button("⬇️ Download Excel", data=df_to_xlsx(edited), file_name="converted_data.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            elif name.endswith(".xls"):
                df = pd.read_excel(uploaded, engine="xlrd")
                edited = st.data_editor(df, num_rows="dynamic", use_container_width=True, key="xls_editor")
                st.download_button("⬇️ Download Excel", data=df_to_xlsx(edited), file_name="converted_data.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            elif name.endswith(".csv"):
                df = pd.read_csv(uploaded)
                edited = st.data_editor(df, num_rows="dynamic", use_container_width=True, key="csv_editor")
                st.download_button("⬇️ Download Excel", data=df_to_xlsx(edited), file_name="converted_data.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            else:
                reader = PdfReader(uploaded)
                pages = [p.extract_text() or "" for p in reader.pages]
                extracted = "\n\n".join(f"--- Page {i+1} ---\n{t}" for i, t in enumerate(pages))
                if extracted.strip():
                    st.text_area("PDF থেকে পাওয়া লেখা (প্রয়োজনে কপি/এডিট করুন)", value=extracted, height=280, key="pdf_text")
                    lines = [line.strip() for line in extracted.splitlines() if line.strip() and not line.startswith("--- Page")]
                    df = pd.DataFrame({"Extracted Text": lines})
                    st.dataframe(df, use_container_width=True)
                    st.download_button("⬇️ Download PDF text as Excel", data=df_to_xlsx(df, "PDF Text"), file_name="pdf_text.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                else:
                    st.error("PDF-এ পড়ার মতো text পাওয়া যায়নি। এটি scanned/image PDF হলে এই সংস্করণে OCR নেই।")
        except Exception as e:
            st.error(f"ফাইলটি পড়া যায়নি: {e}")
            st.caption("সমর্থিত ফাইল ব্যবহার করুন এবং ফাইলটি ক্ষতিগ্রস্ত কি না পরীক্ষা করুন।")

elif page == "👨‍💼 Employee Details":
    st.subheader("👨‍💼 Employee Details Template")
    st.write("নিজের কর্মীদের তথ্য লিখে Excel ডাউনলোড করুন।")
    starter = pd.DataFrame([
        {"Employee ID":"EMP001","Employee Name":"","Department":"","Designation":"","Phone":"","Joining Date":""}
    ])
    df = st.data_editor(starter, num_rows="dynamic", use_container_width=True, key="employee_editor")
    st.download_button("⬇️ Download Employee Excel", data=df_to_xlsx(df, "Employee Details"), file_name="employee_details.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

elif page == "💰 Salary Calculator":
    st.subheader("💰 Salary Calculator")
    st.caption("Basic Salary, HRA, allowance ও deduction লিখুন। Gross/Net salary স্বয়ংক্রিয়ভাবে হিসাব হবে।")
    starter = pd.DataFrame([
        {"Employee ID":"EMP001","Employee Name":"","Basic Salary":15000.0,"HRA":3000.0,"Other Allowance":1000.0,"Deduction":500.0}
    ])
    df = st.data_editor(starter, num_rows="dynamic", use_container_width=True, key="salary_editor",
                        column_config={
                            "Basic Salary": st.column_config.NumberColumn(min_value=0, format="₹ %.2f"),
                            "HRA": st.column_config.NumberColumn(min_value=0, format="₹ %.2f"),
                            "Other Allowance": st.column_config.NumberColumn(min_value=0, format="₹ %.2f"),
                            "Deduction": st.column_config.NumberColumn(min_value=0, format="₹ %.2f")
                        })
    calc = df.copy()
    for col in ["Basic Salary","HRA","Other Allowance","Deduction"]:
        calc[col] = pd.to_numeric(calc[col], errors="coerce").fillna(0)
    calc["Gross Salary"] = calc["Basic Salary"] + calc["HRA"] + calc["Other Allowance"]
    calc["Net Salary"] = calc["Gross Salary"] - calc["Deduction"]
    st.dataframe(calc, use_container_width=True)
    st.download_button("⬇️ Download Salary Excel", data=df_to_xlsx(calc, "Salary Calculation"), file_name="salary_calculation.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

else:
    st.subheader("🗓️ Attendance Sheet")
    st.caption("মাস, কাজের দিন ও উপস্থিতির দিন লিখুন।")
    starter = pd.DataFrame([
        {"Employee ID":"EMP001","Employee Name":"","Month":date.today().strftime("%B %Y"),"Working Days":26,"Present Days":25}
    ])
    df = st.data_editor(starter, num_rows="dynamic", use_container_width=True, key="attendance_editor")
    calc = df.copy()
    calc["Working Days"] = pd.to_numeric(calc["Working Days"], errors="coerce").fillna(0)
    calc["Present Days"] = pd.to_numeric(calc["Present Days"], errors="coerce").fillna(0)
    calc["Absent Days"] = (calc["Working Days"] - calc["Present Days"]).clip(lower=0)
    calc["Attendance %"] = calc.apply(lambda r: (r["Present Days"] / r["Working Days"] * 100) if r["Working Days"] else 0, axis=1)
    st.dataframe(calc, use_container_width=True)
    st.download_button("⬇️ Download Attendance Excel", data=df_to_xlsx(calc, "Attendance"), file_name="attendance_sheet.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

st.divider()
st.caption("AI Excel Generator • Prototype • ফাইলের তথ্য ডাউনলোডের আগে যাচাই করুন।")